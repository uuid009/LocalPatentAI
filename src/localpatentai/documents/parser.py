"""结构化文档解析器（支持 PDF、DOCX、TXT、Markdown），保留页码、章节路径、段落/表格/图注/公式标记。"""

from __future__ import annotations

import re
from io import BytesIO

import fitz
from docx import Document as DocxDocument

from localpatentai.core.errors import AppError
from localpatentai.documents.normalization import decode_text_bytes, normalize_text
from localpatentai.knowledge.models import ParsedSectionBlock

HEADING_LINE_RE = re.compile(
    r"^(?:#{1,6}\s+.+|第[一二三四五六七八九十百0-9]+[章节部分条]\s*.+|"
    r"(?:[一二三四五六七八九十]+|[0-9]+(?:\.[0-9]+){0,2})[、.．\s]\s*\S.+)$"
)
CLAIM_LINE_RE = re.compile(r"^(?:权利要求\s*\d+|【权利要求\s*\d+】|\d+[.、．]\s*一种.+)")
STEP_LINE_RE = re.compile(r"^(?:步骤\s*[S0-9一二三四五六七八九十]+|S\d+[:：.\s]).+", re.IGNORECASE)
FIGURE_LINE_RE = re.compile(r"^(?:图\s*\d+|附图\s*\d+|Fig\.?\s*\d+)[:：.\s—\-]*.*", re.IGNORECASE)
FORMULA_LINE_RE = re.compile(r"(?:\\frac|\\sum|∑|∫|√|[A-Za-z_]+\s*=\s*[^=\n]{3,}\s*\(\d+\))")


def parse_document_bytes(
    *,
    file_ext_type: str,
    content: bytes,
    filename: str,
    parser_version: str = "v1.0",
) -> list[ParsedSectionBlock]:
    if file_ext_type == "pdf":
        return _parse_pdf(content, filename, parser_version)
    if file_ext_type == "docx":
        return _parse_docx(content, filename, parser_version)
    if file_ext_type in {"txt", "md"}:
        return _parse_text_or_markdown(content, filename, parser_version)
    raise AppError(400, "unsupported_parser_type", f"不支持解析的文件类型: {file_ext_type}")


def _classify_line(line: str) -> tuple[str, str | None, str | None]:
    stripped = line.strip()
    if CLAIM_LINE_RE.match(stripped):
        return "claim", None, None
    if FIGURE_LINE_RE.match(stripped):
        fig_match = re.search(r"\d+", stripped)
        fig_id = f"fig-{fig_match.group(0)}" if fig_match else "fig-1"
        return "figure_caption", None, fig_id
    if STEP_LINE_RE.match(stripped):
        return "experiment_step", None, None
    if FORMULA_LINE_RE.search(stripped):
        return "formula", None, None
    if HEADING_LINE_RE.match(stripped) and len(stripped) <= 80:
        return "heading", None, None
    return "paragraph", None, None


def _parse_pdf(content: bytes, filename: str, parser_version: str) -> list[ParsedSectionBlock]:
    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise AppError(422, "pdf_parse_failed", "PDF 文件损坏或无法打开") from exc

    blocks: list[ParsedSectionBlock] = []
    current_section_path: list[str] = [filename]
    para_counter = 0

    try:
        if doc.page_count == 0:
            raise AppError(422, "pdf_empty_pages", "PDF 文档不包含任何页面")
        for page_idx in range(doc.page_count):
            page_num = page_idx + 1
            page = doc[page_idx]
            raw_page_text = page.get_text("text") or ""
            cleaned_page, flags = normalize_text(raw_page_text)
            if not cleaned_page:
                continue
            paragraphs = [p.strip() for p in re.split(r"\n\s*\n", cleaned_page) if p.strip()]
            if len(paragraphs) == 1 and "\n" in paragraphs[0]:
                paragraphs = [line.strip() for line in paragraphs[0].splitlines() if line.strip()]

            for para in paragraphs:
                para_counter += 1
                b_type, table_id, fig_id = _classify_line(para)
                if b_type == "heading":
                    heading_title = re.sub(r"^#+\s*", "", para).strip()
                    current_section_path = [filename, heading_title]
                blocks.append(
                    ParsedSectionBlock(
                        block_type=b_type,  # type: ignore[arg-type]
                        text=para,
                        page_start=page_num,
                        page_end=page_num,
                        section_path=list(current_section_path),
                        paragraph_id=f"p{page_num}-{para_counter}",
                        table_id=table_id,
                        figure_id=fig_id,
                        quality_flags=list(flags),
                        metadata={"parser_version": parser_version},
                    )
                )
    finally:
        doc.close()

    if not blocks:
        raise AppError(
            422,
            "pdf_no_extractable_text",
            "PDF 未能提取到有效文本（可能为纯扫描图像且未启用本地 OCR）",
        )
    return blocks


def _parse_docx(content: bytes, filename: str, parser_version: str) -> list[ParsedSectionBlock]:
    try:
        doc = DocxDocument(BytesIO(content))
    except Exception as exc:
        raise AppError(422, "docx_parse_failed", "DOCX 文件损坏或无法解析") from exc

    blocks: list[ParsedSectionBlock] = []
    current_section_path: list[str] = [filename]
    para_counter = 0

    for paragraph in doc.paragraphs:
        raw = (paragraph.text or "").strip()
        if not raw:
            continue
        cleaned, flags = normalize_text(raw)
        if not cleaned:
            continue
        para_counter += 1
        style_name = (paragraph.style.name if paragraph.style else "").lower()
        if style_name.startswith(("heading", "title")):
            current_section_path = [filename, cleaned]
            b_type = "heading"
            fig_id = None
        else:
            b_type, _, fig_id = _classify_line(cleaned)
            if b_type == "heading":
                current_section_path = [filename, cleaned]

        blocks.append(
            ParsedSectionBlock(
                block_type=b_type,  # type: ignore[arg-type]
                text=cleaned,
                page_start=1,
                page_end=1,
                section_path=list(current_section_path),
                paragraph_id=f"p1-{para_counter}",
                figure_id=fig_id,
                quality_flags=list(flags),
                metadata={"parser_version": parser_version},
            )
        )

    for tbl_idx, table in enumerate(doc.tables, start=1):
        rows_text: list[str] = []
        for row in table.rows:
            cells = [
                normalize_text(cell.text or "")[0].replace("\n", " ").strip() for cell in row.cells
            ]
            if any(cells):
                rows_text.append(" | ".join(cells))
        if rows_text:
            table_id = f"tbl-{tbl_idx}"
            blocks.append(
                ParsedSectionBlock(
                    block_type="table",
                    text="\n".join(rows_text),
                    page_start=1,
                    page_end=1,
                    section_path=list(current_section_path),
                    paragraph_id=f"t1-{tbl_idx}",
                    table_id=table_id,
                    metadata={"parser_version": parser_version, "row_count": len(rows_text)},
                )
            )

    if not blocks:
        raise AppError(422, "docx_empty", "DOCX 文档中未提取到任何正文或表格内容")
    return blocks


def _parse_text_or_markdown(
    content: bytes, filename: str, parser_version: str
) -> list[ParsedSectionBlock]:
    raw_text = decode_text_bytes(content)
    cleaned_doc, doc_flags = normalize_text(raw_text)
    if not cleaned_doc:
        raise AppError(422, "text_empty", "文档内容为空")

    blocks: list[ParsedSectionBlock] = []
    current_section_path: list[str] = [filename]
    lines = cleaned_doc.splitlines()

    para_buffer: list[str] = []
    table_buffer: list[str] = []
    para_counter = 0
    table_counter = 0

    def flush_para() -> None:
        nonlocal para_buffer, para_counter
        if not para_buffer:
            return
        joined = "\n".join(para_buffer).strip()
        para_buffer = []
        if not joined:
            return
        para_counter += 1
        b_type, _, fig_id = _classify_line(joined)
        blocks.append(
            ParsedSectionBlock(
                block_type=b_type,  # type: ignore[arg-type]
                text=joined,
                page_start=1,
                page_end=1,
                section_path=list(current_section_path),
                paragraph_id=f"p1-{para_counter}",
                figure_id=fig_id,
                quality_flags=list(doc_flags),
                metadata={"parser_version": parser_version},
            )
        )

    def flush_table() -> None:
        nonlocal table_buffer, table_counter
        if not table_buffer:
            return
        joined = "\n".join(table_buffer).strip()
        table_buffer = []
        if not joined:
            return
        table_counter += 1
        blocks.append(
            ParsedSectionBlock(
                block_type="table",
                text=joined,
                page_start=1,
                page_end=1,
                section_path=list(current_section_path),
                paragraph_id=f"t1-{table_counter}",
                table_id=f"tbl-{table_counter}",
                quality_flags=list(doc_flags),
                metadata={"parser_version": parser_version},
            )
        )

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush_table()
            flush_para()
            continue

        # Markdown 表格识别
        if stripped.startswith("|") and stripped.endswith("|"):
            flush_para()
            table_buffer.append(stripped)
            continue
        flush_table()

        b_type, _, fig_id = _classify_line(stripped)
        if b_type == "heading":
            flush_para()
            heading_title = re.sub(r"^#+\s*", "", stripped).strip()
            current_section_path = [filename, heading_title]
            para_counter += 1
            blocks.append(
                ParsedSectionBlock(
                    block_type="heading",
                    text=heading_title,
                    page_start=1,
                    page_end=1,
                    section_path=list(current_section_path),
                    paragraph_id=f"p1-{para_counter}",
                    quality_flags=list(doc_flags),
                    metadata={"parser_version": parser_version},
                )
            )
        elif b_type in {"claim", "experiment_step", "figure_caption", "formula"}:
            flush_para()
            para_counter += 1
            blocks.append(
                ParsedSectionBlock(
                    block_type=b_type,  # type: ignore[arg-type]
                    text=stripped,
                    page_start=1,
                    page_end=1,
                    section_path=list(current_section_path),
                    paragraph_id=f"p1-{para_counter}",
                    figure_id=fig_id,
                    quality_flags=list(doc_flags),
                    metadata={"parser_version": parser_version},
                )
            )
        else:
            para_buffer.append(stripped)

    flush_table()
    flush_para()
    return blocks
