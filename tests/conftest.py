"""pytest 全局测试夹具：数据库回滚/清理、三组用户与项目初始化、测试文档生成器。"""

from __future__ import annotations

import uuid
from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import text

from localpatentai.auth.service import (
    _create_session_token,
    register_user_in_group,
)
from localpatentai.core.config import Settings, get_settings
from localpatentai.core.security import hash_password
from localpatentai.db.models import KnowledgeSpace, User
from localpatentai.db.repository import seed_initial_groups_and_spaces
from localpatentai.db.session import get_session_factory
from localpatentai.main import app
from localpatentai.projects.service import create_project
from localpatentai.workflows.graph import close_workflow_checkpointer


@pytest.fixture()
def test_settings(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Settings:
    """使用测试专用本地存储目录和确定性 mock 嵌入/重排/LLM 后端。"""
    storage_dir = tmp_path / "storage"
    uploads_dir = storage_dir / "uploads"
    artifacts_dir = storage_dir / "artifacts"
    workspaces_dir = storage_dir / "workspaces"
    uploads_dir.mkdir(parents=True, exist_ok=True)
    artifacts_dir.mkdir(parents=True, exist_ok=True)
    workspaces_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setenv("DATA_ROOT", str(workspaces_dir))
    monkeypatch.setenv("UPLOAD_DIR", str(uploads_dir))
    monkeypatch.setenv("ARTIFACT_DIR", str(artifacts_dir))
    monkeypatch.setenv("EMBEDDING_BACKEND", "mock")
    monkeypatch.setenv("RERANKER_BACKEND", "mock")
    monkeypatch.setenv("LLM_PROVIDER", "mock")
    get_settings.cache_clear()
    settings = get_settings()
    yield settings
    get_settings.cache_clear()


@pytest.fixture()
def db_session(test_settings: Settings):
    """每次测试前清理测试数据并重新初始化三组（algorithm / framework / cloud）与公共空间。"""
    close_workflow_checkpointer()
    session_factory = get_session_factory()
    with session_factory() as session:
        session.execute(
            text(
                "TRUNCATE TABLE audit_events, artifacts, tasks, chunks, document_versions, "
                "documents, project_members, projects, knowledge_spaces, group_members, "
                "auth_sessions, users, groups RESTART IDENTITY CASCADE"
            )
        )
        session.commit()
        seed_initial_groups_and_spaces(session)
        session.commit()
        yield session
        session.rollback()
    close_workflow_checkpointer()


@pytest.fixture()
def seeded_org(db_session, test_settings: Settings) -> dict:
    """
    初始化完整三组组织环境（规格 12.1 节）：
    - user_algo (algorithm 组)
    - user_algo_member (algorithm 组普通成员)
    - user_framework (framework 组)
    - user_cloud (cloud 组)
    - user_admin (system_admin，无业务组成员身份)
    - project_algo, project_framework, project_cloud
    """
    u_algo, _, token_algo = register_user_in_group(
        db_session,
        username="user_algo",
        password="Password123!",
        display_name="张算法",
        group_code="algorithm",
        group_role="group_admin",
    )
    u_algo_member, _, token_algo_member = register_user_in_group(
        db_session,
        username="user_algo_member",
        password="Password123!",
        display_name="李算法成员",
        group_code="algorithm",
        group_role="member",
    )
    u_framework, _, token_framework = register_user_in_group(
        db_session,
        username="user_framework",
        password="Password123!",
        display_name="王框架",
        group_code="framework",
        group_role="group_admin",
    )
    u_cloud, _, token_cloud = register_user_in_group(
        db_session,
        username="user_cloud",
        password="Password123!",
        display_name="赵云原生",
        group_code="cloud",
        group_role="group_admin",
    )

    # 创建不隶属任何科研业务组的 system_admin
    u_admin = User(
        username="sys_admin",
        display_name="系统管理员",
        password_hash=hash_password("Password123!"),
        system_role="system_admin",
        is_active=True,
    )
    db_session.add(u_admin)
    db_session.flush()
    token_admin = _create_session_token(db_session, u_admin.id)
    db_session.commit()

    p_algo, ks_algo_priv = create_project(
        db_session,
        user_id=u_algo.id,
        group_code="algorithm",
        name="算法组混合精度稀疏注意力专利项目",
        description="面向端侧大模型推理的动态块级 FP8/INT4 混合精度稀疏注意力方案",
    )
    p_framework, ks_framework_priv = create_project(
        db_session,
        user_id=u_framework.id,
        group_code="framework",
        name="框架组异构计算图算子融合专利项目",
        description="静态内存规划与跨流算子融合调度器",
    )
    p_cloud, ks_cloud_priv = create_project(
        db_session,
        user_id=u_cloud.id,
        group_code="cloud",
        name="云组算力容器拓扑感知热迁移专利项目",
        description="基于芯片互联拓扑亲和度评分的弹性容器热迁移方法",
    )

    ks_common = (
        db_session.query(KnowledgeSpace)
        .filter(KnowledgeSpace.scope_type == "approved_common")
        .one()
    )
    ks_algo_shared = (
        db_session.query(KnowledgeSpace)
        .filter(
            KnowledgeSpace.scope_type == "group_shared",
            KnowledgeSpace.group_id == p_algo.group_id,
        )
        .one()
    )

    return {
        "users": {
            "algo": u_algo,
            "algo_member": u_algo_member,
            "framework": u_framework,
            "cloud": u_cloud,
            "admin": u_admin,
        },
        "tokens": {
            "algo": token_algo,
            "algo_member": token_algo_member,
            "framework": token_framework,
            "cloud": token_cloud,
            "admin": token_admin,
        },
        "projects": {
            "algo": p_algo,
            "framework": p_framework,
            "cloud": p_cloud,
        },
        "spaces": {
            "algo_private": ks_algo_priv,
            "framework_private": ks_framework_priv,
            "cloud_private": ks_cloud_priv,
            "algo_shared": ks_algo_shared,
            "common": ks_common,
        },
    }


@pytest.fixture()
def client(test_settings: Settings, db_session) -> TestClient:
    return TestClient(app)


def make_sample_pdf_bytes(title: str, body_text: str) -> bytes:
    """生成带中文标题和正文的 PDF 字节流用于上传解析测试。"""
    doc = fitz.open()
    page = doc.new_page()
    html = f"<h1>{title}</h1><p>{body_text}</p>"
    page.insert_htmlbox(page.rect + (36, 36, -36, -36), html)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def make_sample_docx_bytes(tmp_path: Path, title: str, paragraphs: list[str]) -> bytes:
    """生成带标题、段落与表格的 DOCX 字节流。"""
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标名称"
    table.cell(0, 1).text = "测试数值"
    table.cell(1, 0).text = "首字延迟 TTFT"
    table.cell(1, 1).text = "59.2 ms"
    out = tmp_path / f"{uuid.uuid4().hex}.docx"
    doc.save(str(out))
    return out.read_bytes()
